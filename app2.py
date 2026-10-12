"""
Script de automação de planilha Excel com openpyxl
Autor: Rodrigo
"""

import openpyxl

# ================================
# PARTE 1: CRIAR a planilha
# ================================
print(">>> Criando a planilha...")

book = openpyxl.Workbook()
print("Planilhas existentes:", book.sheetnames)

# Criar a aba "Frutas"
book.create_sheet("Frutas")

# Selecionar a aba
frutas_page = book["Frutas"]

# Adicionar os dados
frutas_page.append(["Fruta", "Quantidade", "Preço"])
frutas_page.append(["Banana", "5", "R$ 3,90"])
frutas_page.append(["Mamão", "2", "R$ 5,90"])
frutas_page.append(["Kiwi", "5", "R$ 20,90"])
frutas_page.append(["Melancia", "1", "R$ 11,90"])

# Salvar a planilha
book.save("Planilha de Compras.xlsx")
print(">>> Planilha criada com sucesso!\n")


# ================================
# PARTE 2: LER a planilha (simples)
# ================================
print(">>> Lendo os dados (simples)...")

book = openpyxl.load_workbook("Planilha de Compras.xlsx")
frutas_page = book["Frutas"]

for rows in frutas_page.iter_rows(min_row=2, max_row=5):
    for cell in rows:
        print(cell.value)

print()


# ================================
# PARTE 3: LER a planilha (formatado)
# ================================
print(">>> Lendo os dados (formatado)...")

for rows in frutas_page.iter_rows(min_row=2, max_row=5):
    print(f"{rows[0].value}, {rows[1].value}, {rows[2].value}")

print()


# ================================
# PARTE 4: MODIFICAR e SALVAR
# ================================
print(">>> Modificando 'Banana' para 'Fruta 1'...")

for rows in frutas_page.iter_rows(min_row=2, max_row=5):
    for cell in rows:
        if cell.value == "Banana":
            cell.value = "Fruta 1"

book.save("Planilha de Compras.xlsx")
print(">>> Planilha atualizada com sucesso!")


# ================================
# PARTE 5: CONFERIR se mudou
# ================================
print("\n>>> Conferindo o resultado final...")

book = openpyxl.load_workbook("Planilha de Compras.xlsx")
frutas_page = book["Frutas"]

for rows in frutas_page.iter_rows(min_row=2, max_row=5):
    print(f"{rows[0].value}, {rows[1].value}, {rows[2].value}")