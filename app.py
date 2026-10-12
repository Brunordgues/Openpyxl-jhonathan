import openpyxl

#criar uma planilha do excel(book)
book = openpyxl.Workbook()

#como visualizar as planilhas do excel
print(book.sheetnames)

#como criar uma página
book.create_sheet("Frutas")

#como selecionar uma página
frutas_page = book["Frutas"]
frutas_page.append(["Fruta", "Quantidade", "Preço"])
frutas_page.append(["Banana", "5", "R$ 3,90"])
frutas_page.append(["Mamão", "2", "R$ 5,90"])
frutas_page.append(["Kiwi", "5", "R$ 20,90"])
frutas_page.append(["Melancia", "1", "R$ 11,90"])

#salvar a planilha
book.save("Planilha de Compras.xlsx")